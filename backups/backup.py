import dearpygui.dearpygui as dpg
from openpyxl.utils.cell import column_index_from_string
import json
import ast

import column_extraction

dpg.create_context()


def range_to_string(input):
    '''
    "range" : {
        "type": "End of column"
    },
    "range" : {
        "type": "Up to row",
        "row": n
    },
    "range" : {
        "type": "Up to code",
        "code": "stop"
        "column": "A"
    }

    End of column,
    To row: n,
    To code "stop" at "A"
    '''
    
    if input["type"] == "End of column":
        return "End of column"
    elif input["type"] == "Up to row":
        return f"To row: {input["row"]}"
    elif input["type"] == "Up to code":
        return f"To code \"{input["code"]}\" at \"{input["column"]}\""
    

def string_to_range(input):
    if input == "End of column":
        return {"type": "End of column"}
    elif input[0:8] == "To row: ":
        return {
            "type": "Up to row",
            "row": int(input[8:])
        }
    elif input[0:8] == "To code ":
        first_q = 8
        i = first_q +1
        while input[i] != "\"":
            i += 1
        second_q = i
        third_q = i + 6
        last_q = len(input)-1
        return {
            "type": "Up to code",
            "code": input[first_q+1:second_q],
            "column": input[third_q:last_q]
        }


def add_letter():
    letter = dpg.get_value("Column Input").upper()

    try:
        column_index_from_string(letter)
    except:
        return

    if letter:
        item_tag = f"item_{len(dpg.get_item_children('Columns List', 1))}"
        with dpg.group(parent="Columns List", horizontal=True, tag=item_tag):
            dpg.add_text(letter, label=letter)
            dpg.add_button(label="-", callback=lambda s, a, u: dpg.delete_item(item_tag))


def save_to_file():
    entries = []

    rows = dpg.get_item_children("table", 1)
    for row in rows:
        cells = dpg.get_item_children(row, 1)

        entry = {}
        entry["name"] = dpg.get_item_configuration(cells[0])["label"]
        entry["enabled"] = True if dpg.get_item_configuration(cells[1])["label"] == "True" else False
        entry["input"] = {}
        entry["input"]["filepath"] = dpg.get_item_configuration(cells[2])["label"]
        entry["input"]["sheetname"] = dpg.get_item_configuration(cells[3])["label"]
        entry["input"]["row_offset"] = int(dpg.get_item_configuration(cells[4])["label"])
        entry["input"]["columns"] = ast.literal_eval(dpg.get_item_configuration(cells[5])["label"])
        entry["input"]["range_mode"] = string_to_range(dpg.get_item_configuration(cells[6])["label"])
        entry["output"] = {}
        entry["output"]["filepath"] = dpg.get_item_configuration(cells[7])["label"]
        entry["output"]["sheetname"] = dpg.get_item_configuration(cells[8])["label"]
        entry["output"]["insert_mode"] = dpg.get_item_configuration(cells[9])["label"]
        entry["output"]["include_name"] = True if dpg.get_item_configuration(cells[10])["label"] == "True" else False

        print(entry)
        entries.append(entry)
    
    with open('project_info.json', 'w') as json_file:
        json.dump(entries, json_file, indent=2)
    
        
def load_to_table(filepath: str):
    with open('project_info.json', 'r') as json_file:
        entries = json.load(json_file)
        for entry in entries:
            name = entry["name"]
            enabled = entry["enabled"]

            input_fp = entry["input"]["filepath"]
            input_st = entry["input"]["sheetname"]
            row_offset = entry["input"]["row_offset"]
            columns = entry["input"]["columns"]
            range_mode = range_to_string(entry["input"]["range_mode"])

            output_fp = entry["output"]["filepath"]
            output_st = entry["output"]["sheetname"]
            insert_mode = entry["output"]["insert_mode"]
            include_name = entry["output"]["include_name"]

            with dpg.table_row(parent="table"):
                dpg.add_selectable(label=name, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=enabled, span_columns=True, callback=table_row_callback, user_data=entry)

                dpg.add_selectable(label=input_fp, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=input_st, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=row_offset, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=columns, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=range_mode, span_columns=True, callback=table_row_callback, user_data=entry)

                dpg.add_selectable(label=output_fp, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=output_st, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=insert_mode, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=include_name, span_columns=True, callback=table_row_callback, user_data=entry)

    json_file.close()


def save_to_table():
    name =  dpg.get_value("Name Input")
    if name == "":
        return
    rows = dpg.get_item_children("table", 1)
    for row in rows:
        cells = dpg.get_item_children(row, 1)
        if dpg.get_item_label(cells[0]) == name:
            return
    
    enabled = dpg.get_value("Enabled Input")

    input_fp = dpg.get_value("Input FP Input")
    input_st = dpg.get_value("Input ST Input")
    row_offset = dpg.get_value("Row Offset Input")

    columns = []
    column_children = dpg.get_item_children("Columns List", 1)
    for i in range(0, len(column_children)):
        child = column_children[i]
        print(dpg.get_item_configuration(dpg.get_item_children(child, 1)[1]))
        columns.append(dpg.get_item_configuration(dpg.get_item_children(child, 1)[0])["label"])

    range_mode = ""
    print(dpg.get_value("Range Input"))
    if dpg.get_value("Range Input") == "End Of Column":
        range_mode = range_to_string({
            "type": "End of column"
        })
    elif dpg.get_value("Range Input") == "Up to the nth row":
        range_mode = range_to_string({
            "type": "Up to row",
            "row": dpg.get_value("row option")
        })
    elif dpg.get_value("Range Input") == "Up to the code":
        try:
            column_index_from_string(dpg.get_value("code option 1"))
        except:
            return
        range_mode = range_to_string({
            "type": "Up to code",
            "code": dpg.get_value("code option 0"),
            "column": dpg.get_value("code option 1")
        })

    output_fp = dpg.get_value("Output FP Input")
    output_st = dpg.get_value("Output ST Input")
    insert_mode = dpg.get_value("Insert Mode Input")
    include_name = dpg.get_value("Include Name Input")

    entry = {
        "name": name,
        "enabled": enabled,
        "input": {
            "filepath": input_fp,
            "sheetname": input_st,
            "row_offset": int(row_offset),
            "columns": columns,
            "range_mode": string_to_range(range_mode)
        },
        "output": {
            "filepath": output_fp,
            "sheetname": output_st,
            "insert_mode": insert_mode,
            "include_name": include_name
        }
    }
    print(entry)

    with dpg.table_row(parent="table"):
        dpg.add_selectable(label=name, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=enabled, span_columns=True, callback=table_row_callback, user_data=entry)

        dpg.add_selectable(label=input_fp, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=input_st, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=row_offset, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=columns, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=range_mode, span_columns=True, callback=table_row_callback, user_data=entry)

        dpg.add_selectable(label=output_fp, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=output_st, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=insert_mode, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=include_name, span_columns=True, callback=table_row_callback, user_data=entry)


def delete_entry():
    name =  dpg.get_value("Name Input")

    rows = dpg.get_item_children("table", 1)
    for row in rows:
        cells = dpg.get_item_children(row, 1)
        if dpg.get_item_label(cells[0]) != name:
            continue
        dpg.delete_item(row)
        break


def update_entry():
    name = dpg.get_value("Name Input")
    if name == "":
        return
        
    rows = dpg.get_item_children("table", 1)
    for row in rows:
        cells = dpg.get_item_children(row, 1)
        if dpg.get_item_label(cells[0]) != name:
            continue

        enabled = dpg.get_value("Enabled Input")

        input_fp = dpg.get_value("Input FP Input")
        input_st = dpg.get_value("Input ST Input")
        row_offset = dpg.get_value("Row Offset Input")

        columns = []
        column_children = dpg.get_item_children("Columns List", 1)
        for i in range(0, len(column_children)):
            child = column_children[i]
            columns.append(dpg.get_item_configuration(dpg.get_item_children(child, 1)[0])["label"])

# "End Of Column", "Up to the nth row", "Up to the code"
        range_mode = ""
        print(dpg.get_value("Range Input"))
        if dpg.get_value("Range Input") == "End Of Column":
            range_mode = range_to_string({
                "type": "End of column"
            })
        elif dpg.get_value("Range Input") == "Up to the nth row":
            range_mode = range_to_string({
                "type": "Up to row",
                "row": dpg.get_value("row option")
            })
        elif dpg.get_value("Range Input") == "Up to the code":
            range_mode = range_to_string({
                "type": "Up to code",
                "code": dpg.get_value("code option 0"),
                "column": dpg.get_value("code option 1")
            })

        output_fp = dpg.get_value("Output FP Input")
        output_st = dpg.get_value("Output ST Input")
        insert_mode = dpg.get_value("Insert Mode Input")
        include_name = dpg.get_value("Include Name Input")

        entry = {
            "name": name,
            "enabled": enabled,
            "input": {
                "filepath": input_fp,
                "sheetname": input_st,
                "row_offset": int(row_offset),
                "columns": columns,
                "range_mode": string_to_range(range_mode)
            },
            "output": {
                "filepath": output_fp,
                "sheetname": output_st,
                "insert_mode": insert_mode,
                "include_name": include_name
            }
        }

        dpg.set_item_user_data(cells[0], entry)

        dpg.set_item_label(cells[1], enabled)
        dpg.set_item_user_data(cells[1], entry)
        
        dpg.set_item_label(cells[2], input_fp)
        dpg.set_item_user_data(cells[2], entry)

        dpg.set_item_label(cells[3], input_st)
        dpg.set_item_user_data(cells[3], entry)

        dpg.set_item_label(cells[4], row_offset)
        dpg.set_item_user_data(cells[4], entry)

        dpg.set_item_label(cells[5], columns)
        dpg.set_item_user_data(cells[5], entry)

        dpg.set_item_label(cells[6], range_mode)
        dpg.set_item_user_data(cells[6], entry)

        dpg.set_item_label(cells[7], output_fp)
        dpg.set_item_user_data(cells[7], entry)

        dpg.set_item_label(cells[8], output_st)
        dpg.set_item_user_data(cells[8], entry)

        dpg.set_item_label(cells[9], insert_mode)
        dpg.set_item_user_data(cells[9], entry)

        dpg.set_item_label(cells[10], include_name)
        dpg.set_item_user_data(cells[10], entry)

        break


def table_row_callback(sender, app_data, user_data):
    name = user_data["name"]
    enabled = user_data["enabled"]

    input_fp = user_data["input"]["filepath"]
    input_st = user_data["input"]["sheetname"]
    row_offset = user_data["input"]["row_offset"]
    columns = user_data["input"]["columns"]
    range_mode = user_data["input"]["range_mode"]
    print(range_mode)

    output_fp = user_data["output"]["filepath"]
    output_st = user_data["output"]["sheetname"]
    insert_mode = user_data["output"]["insert_mode"]
    include_name = user_data["output"]["include_name"]
    
    dpg.set_value("Name Input", name)
    dpg.set_value("Enabled Input", enabled)
    dpg.set_value("Input FP Input", input_fp)
    dpg.set_value("Input ST Input", input_st)
    dpg.set_value("Row Offset Input", row_offset)
    
    column_children = dpg.get_item_children("Columns List", 1)
    for child in column_children:
        dpg.delete_item(child)
    for col in columns:
        dpg.set_value("Column Input", col)
        add_letter()
    dpg.set_value("Column Input", "")

    print(range_mode["type"])
    if range_mode["type"] == "End of column":
        dpg.set_value("Range Input", "End Of Column")
        dpg.hide_item("row option")
        dpg.hide_item("code option 0")
        dpg.hide_item("code option 1")
    elif range_mode["type"] == "Up to row":
        dpg.set_value("Range Input", "Up to the nth row")
        dpg.set_value("row option", range_mode["row"])
        dpg.show_item("row option")
        dpg.hide_item("code option 0")
        dpg.hide_item("code option 1")
    elif range_mode["type"] == "Up to code":
        dpg.set_value("Range Input", "Up to the code")
        dpg.set_value("code option 0", range_mode["code"])
        dpg.set_value("code option 1", range_mode["column"])
        dpg.hide_item("row option")
        dpg.show_item("code option 0")
        dpg.show_item("code option 1")
    
    dpg.set_value("Output FP Input", output_fp)
    dpg.set_value("Output ST Input", output_st)
    dpg.set_value("Insert Mode Input", insert_mode)
    dpg.set_value("Include Name Input", include_name)


def radio_callback(sender, app_data, user_data):
    selected_value = app_data
    if selected_value == "End Of Column":
        dpg.hide_item("row option")
        dpg.hide_item("code option 0")
        dpg.hide_item("code option 1")
    elif selected_value == "Up to the nth row":
        dpg.show_item("row option")
        dpg.hide_item("code option 0")
        dpg.hide_item("code option 1")
    elif selected_value == "Up to the code":
        dpg.hide_item("row option")
        dpg.show_item("code option 0")
        dpg.show_item("code option 1")


def create_file_input(input_tag, entry_label, button_label="Browse", dialog_tag=None):
    if dialog_tag is None:
        dialog_tag = f"{input_tag}_dialog"

    # Callback function to handle file selection
    def file_callback(sender, app_data):
        dpg.set_value(input_tag, app_data['file_path_name'])

    
    dpg.add_input_text(tag=input_tag, hint="Select a file", label=entry_label, readonly=True)
    dpg.add_button(label=button_label, callback=lambda: dpg.show_item(dialog_tag))

    with dpg.file_dialog(directory_selector=False, show=False, callback=file_callback, tag=dialog_tag, width=600, height=500):
        dpg.add_file_extension(".xlsx")


def add_error_msg_to_console(message):
    dpg.add_text(message, parent="console_window", color=(196, 43, 43))
    dpg.set_y_scroll("console_window", dpg.get_y_scroll_max("console_window"))


def add_valid_msg_to_console(message):
    dpg.add_text(message, parent="console_window", color=(69, 214, 69))
    dpg.set_y_scroll("console_window", dpg.get_y_scroll_max("console_window"))


def pull_columns():
    stati = column_extraction.run()
    for status in stati:
        if status[1] == "v":
            add_valid_msg_to_console(status)
        elif status[1] == "e":
            add_error_msg_to_console(status)


with dpg.window(tag="Primary Window"):
    with dpg.tab_bar():
        with dpg.tab(label="Columns"):
            with dpg.child_window(tag="Name Window", autosize_x=True, height=75):
                dpg.add_input_text(label="Name", tag="Name Input")
                dpg.add_checkbox(label="Enabled", tag="Enabled Input")
            
            with dpg.child_window(tag="Input File Information", autosize_x=True, height=250):
                create_file_input(entry_label="Input Filepath", input_tag="Input FP Input")
                dpg.add_input_text(label="Input Sheetname", tag="Input ST Input")
                dpg.add_input_int(label="Row Offset", default_value=1, min_value=1, tag="Row Offset Input")

                dpg.add_input_text(label="Column", tag="Column Input")
                dpg.add_button(label="Add", callback=add_letter)
                dpg.add_separator()
                dpg.add_group(tag="Columns List")

                dpg.add_radio_button(items=["End Of Column", "Up to the nth row", "Up to the code"], callback=radio_callback, horizontal=True, tag="Range Input")

                dpg.add_input_int(label="Row", default_value=1, min_value=1, tag="row option")
                dpg.add_input_text(label="Code", tag="code option 0")
                dpg.add_input_text(label="Column", tag="code option 1")
                
                dpg.hide_item("row option")
                dpg.hide_item("code option 0")
                dpg.hide_item("code option 1")

            with dpg.child_window(tag="Output File Information", autosize_x=True, height=150):
                create_file_input(entry_label="Output Filepath", input_tag="Output FP Input")
                dpg.add_input_text(label="Output Sheetname", tag="Output ST Input")
                dpg.add_combo(items=["Next empty row", "Next empty column"], default_value="Next empty row", tag="Insert Mode Input")
                dpg.add_checkbox(label="Include Name", tag="Include Name Input")

            with dpg.group(horizontal=True):
                dpg.add_button(label="Add to table", callback=save_to_table)
                dpg.add_button(label="Update entry", callback=update_entry)
                dpg.add_button(label="Delete entry", callback=delete_entry)
                dpg.add_button(label="Save to file", callback=save_to_file)
                dpg.add_button(label="Pull data", callback=pull_columns)


            with dpg.table(tag="table", header_row=True, row_background=True, policy=dpg.mvTable_SizingStretchSame,
                        borders_innerH=True, borders_outerH=True, borders_innerV=True,
                        borders_outerV=True):
                dpg.add_table_column(label="Name")
                dpg.add_table_column(label="Enabled")
                dpg.add_table_column(label="Input FP")
                dpg.add_table_column(label="Input ST")
                dpg.add_table_column(label="Row Offset")
                dpg.add_table_column(label="Columns")
                dpg.add_table_column(label="Range Mode")
                dpg.add_table_column(label="Output FP")
                dpg.add_table_column(label="Output ST")
                dpg.add_table_column(label="Insert Mode")
                dpg.add_table_column(label="Include Name")
        
        with dpg.tab(label="Rows"):
            with dpg.child_window(autosize_x=True, height=75):
                dpg.add_input_text(label="Name")
                dpg.add_checkbox(label="Enabled")
            
            with dpg.child_window(autosize_x=True, height=250):
                create_file_input(entry_label="Input Filepath", input_tag="##0")
                dpg.add_input_text(label="Input Sheetname")
                dpg.add_input_int(label="Row Offset", default_value=1, min_value=1)

                dpg.add_radio_button(items=["Row", "Code"], callback=radio_callback, horizontal=True)

            with dpg.child_window(autosize_x=True, height=150):
                create_file_input(entry_label="Output Filepath", input_tag="##1")
                dpg.add_input_text(label="Output Sheetname")
                dpg.add_checkbox(label="Include Name")

            with dpg.group(horizontal=True):
                dpg.add_button(label="Add to table", callback=save_to_table)
                dpg.add_button(label="Update entry", callback=update_entry)
                dpg.add_button(label="Delete entry", callback=delete_entry)
                dpg.add_button(label="Save to file", callback=save_to_file)
                dpg.add_button(label="Pull data", callback=pull_columns)

            with dpg.table(tag="row table", header_row=True, row_background=True, policy=dpg.mvTable_SizingStretchSame,
                        borders_innerH=True, borders_outerH=True, borders_innerV=True,
                        borders_outerV=True):
                dpg.add_table_column(label="Name")
                dpg.add_table_column(label="Enabled")
                dpg.add_table_column(label="Input FP")
                dpg.add_table_column(label="Input ST")
                dpg.add_table_column(label="Selection Mode")
                dpg.add_table_column(label="Output FP")
                dpg.add_table_column(label="Output ST")
                dpg.add_table_column(label="Include Name")

    with dpg.child_window(menubar=True, horizontal_scrollbar=True, autosize_x=True, height=500, tag="console_window"):
        with dpg.menu_bar():
            dpg.add_menu(label="Status Reports")

    load_to_table("project_info.json")


dpg.create_viewport(title='Custom Title', width=900, height=900)
dpg.setup_dearpygui()

dpg.show_viewport()
dpg.set_primary_window("Primary Window", True)
dpg.start_dearpygui()


dpg.destroy_context()