from pygui_utils import *
import json
from openpyxl.utils.cell import column_index_from_string, get_column_letter

import row_extraction

def selection_to_string(selection: json):
    '''
    {
        "type": row,
        "row": x
    },
    {
        "type": code,
        "code": "asd",
        "column": "A"
    }
    '''
    if selection["type"] == "row":
        return f"Row {selection["row"]}"
    elif selection["type"] == "code":
        return f"Code \"{selection["code"]}\" at \"{selection["column"]}\""


def string_to_selection(string: str):
    if string[0:3] == "Row":
        return {
            "type": "row",
            "row": int(string[4:])
        }
    if string[0:4] == "Code":
        first_q = 5
        i = first_q +1
        while string[i] != "\"":
            i += 1
        second_q = i
        third_q = i + 6
        last_q = len(string)-1
        return {
            "type": "code",
            "code": string[first_q+1:second_q],
            "column": string[third_q:last_q]
        }


def radio_callback(sender, app_data, user_data):
    selected_value = app_data
    if selected_value == "Row":
        dpg.show_item("Rows row option")
        dpg.hide_item("Rows code option")
        dpg.hide_item("Rows column option")
    elif selected_value == "Code":
        dpg.hide_item("Rows row option")
        dpg.show_item("Rows code option")
        dpg.show_item("Rows column option")


def table_row_callback(sender, app_data, user_data):
    name = user_data["name"]
    enabled = user_data["enabled"]

    input_fp = user_data["input"]["filepath"]
    input_st = user_data["input"]["sheetname"]
    selection_mode = user_data["input"]["selection_mode"]

    output_fp = user_data["output"]["filepath"]
    output_st = user_data["output"]["sheetname"]
    include_name = user_data["output"]["include_name"]
    
    dpg.set_value("Rows Name Input", name)
    dpg.set_value("Rows Enabled Input", enabled)
    dpg.set_value("Rows Input FP Input", input_fp)
    dpg.set_value("Rows Input ST Input", input_st)

    if selection_mode["type"] == "row":
        dpg.set_value("Rows Selection Mode Input", "Row")
        dpg.set_value("Rows row option", selection_mode["row"])
        dpg.show_item("Rows row option")
        dpg.hide_item("Rows code option")
        dpg.hide_item("Rows column option")
    elif selection_mode["type"] == "code":
        dpg.set_value("Rows Selection Mode Input", "Code")
        dpg.set_value("Rows code option", selection_mode["code"])
        dpg.set_value("Rows column option", selection_mode["column"])
        dpg.hide_item("Rows row option")
        dpg.show_item("Rows code option")
        dpg.show_item("Rows column option")
    
    dpg.set_value("Rows Output FP Input", output_fp)
    dpg.set_value("Rows Output ST Input", output_st)
    dpg.set_value("Rows Include Name Input", include_name)


def save_to_row_table():
    name =  dpg.get_value("Rows Name Input")
    if name == "":
        return
    rows = dpg.get_item_children("row table", 1)
    for row in rows:
        cells = dpg.get_item_children(row, 1)
        if dpg.get_item_label(cells[0]) == name:
            return
    
    enabled = dpg.get_value("Rows Enabled Input")

    input_fp = dpg.get_value("Rows Input FP Input")
    input_st = dpg.get_value("Rows Input ST Input")
    selection = dpg.get_value("Rows Selection Mode Input")
    selection_mode = ""
    if selection == "Row":
        selection_mode = f"Row {dpg.get_value("Rows row option")}"
    elif selection == "Code":
        selection_mode = f"Code \"{dpg.get_value("Rows code option")}\" at \"{dpg.get_value("Rows column option")}\""
    try:
        column_index_from_string(dpg.get_value("Rows column option"))
    except:
        return
    
    output_fp = dpg.get_value("Rows Output FP Input")
    output_st = dpg.get_value("Rows Output ST Input")
    include_name = dpg.get_value("Rows Include Name Input")

    entry = {
        "name": name,
        "enabled": enabled,
        "input": {
            "filepath": input_fp,
            "sheetname": input_st,
            "selection_mode": string_to_selection(selection_mode)
        },
        "output": {
            "filepath": output_fp,
            "sheetname": output_st,
            "include_name": include_name
        }
    }
    print(entry)

    with dpg.table_row(parent="row table"):
        dpg.add_selectable(label=name, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=enabled, span_columns=True, callback=table_row_callback, user_data=entry)

        dpg.add_selectable(label=input_fp, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=input_st, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=selection_mode, span_columns=True, callback=table_row_callback, user_data=entry)

        dpg.add_selectable(label=output_fp, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=output_st, span_columns=True, callback=table_row_callback, user_data=entry)
        dpg.add_selectable(label=include_name, span_columns=True, callback=table_row_callback, user_data=entry)


def update_row_entry():
    name = dpg.get_value("Rows Name Input")
    if name == "":
        return
        
    rows = dpg.get_item_children("row table", 1)
    for row in rows:
        cells = dpg.get_item_children(row, 1)
        if dpg.get_item_label(cells[0]) != name:
            continue

        enabled = dpg.get_value("Rows Enabled Input")

        input_fp = dpg.get_value("Rows Input FP Input")
        input_st = dpg.get_value("Rows Input ST Input")
        selection = dpg.get_value("Rows Selection Mode Input")
        selection_mode = ""
        if selection == "Row":
            selection_mode = f"Row {dpg.get_value("Rows row option")}"
        elif selection == "Code":
            selection_mode = f"Code \"{dpg.get_value("Rows code option")}\" at \"{dpg.get_value("Rows column option")}\""
        try:
            column_index_from_string(dpg.get_value("Rows column option"))
        except:
            return

        output_fp = dpg.get_value("Rows Output FP Input")
        output_st = dpg.get_value("Rows Output ST Input")
        include_name = dpg.get_value("Rows Include Name Input")

        entry = {
            "name": name,
            "enabled": enabled,
            "input": {
                "filepath": input_fp,
                "sheetname": input_st,
                "selection_mode": string_to_selection(selection_mode)
            },
            "output": {
                "filepath": output_fp,
                "sheetname": output_st,
                "include_name": include_name
            }
        }

        dpg.set_item_user_data(cells[0], entry)

        dpg.set_item_label(cells[1], enabled)
        dpg.set_item_user_data(cells[1], entry)
        
        dpg.set_item_label(cells[2], input_fp)
        dpg.set_item_user_data(cells[2], entry)

        dpg.set_item_label(cells[3], input_st)
        dpg.set_item_user_data(cells[3], entry)

        dpg.set_item_label(cells[4], selection_mode)
        dpg.set_item_user_data(cells[4], entry)

        dpg.set_item_label(cells[5], output_fp)
        dpg.set_item_user_data(cells[5], entry)

        dpg.set_item_label(cells[6], output_st)
        dpg.set_item_user_data(cells[6], entry)

        dpg.set_item_label(cells[7], include_name)
        dpg.set_item_user_data(cells[7], entry)   


def delete_row_entry():
    name =  dpg.get_value("Rows Name Input")

    rows = dpg.get_item_children("row table", 1)
    for row in rows:
        cells = dpg.get_item_children(row, 1)
        if dpg.get_item_label(cells[0]) != name:
            continue
        dpg.delete_item(row)
        break


def save_to_row_file():
    entries = []

    rows = dpg.get_item_children("row table", 1)
    for row in rows:
        cells = dpg.get_item_children(row, 1)

        entry = {}
        entry["name"] = dpg.get_item_configuration(cells[0])["label"]
        entry["enabled"] = True if dpg.get_item_configuration(cells[1])["label"] == "True" else False
        entry["input"] = {}
        entry["input"]["filepath"] = dpg.get_item_configuration(cells[2])["label"]
        entry["input"]["sheetname"] = dpg.get_item_configuration(cells[3])["label"]
        entry["input"]["selection_mode"] = string_to_selection(dpg.get_item_configuration(cells[4])["label"])
        entry["output"] = {}
        entry["output"]["filepath"] = dpg.get_item_configuration(cells[5])["label"]
        entry["output"]["sheetname"] = dpg.get_item_configuration(cells[6])["label"]
        entry["output"]["include_name"] = True if dpg.get_item_configuration(cells[7])["label"] == "True" else False

        print(entry)
        entries.append(entry)
    
    with open('row_project_info.json', 'w') as json_file:
        json.dump(entries, json_file, indent=2)


def add_error_msg_to_console(message):
    dpg.add_text(message, parent="console_window", color=(196, 43, 43))
    dpg.set_y_scroll("console_window", dpg.get_y_scroll_max("console_window"))


def add_valid_msg_to_console(message):
    dpg.add_text(message, parent="console_window", color=(69, 214, 69))
    dpg.set_y_scroll("console_window", dpg.get_y_scroll_max("console_window"))


def pull_rows():
    stati = row_extraction.run()
    for status in stati:
        if status[1] == "v":
            add_valid_msg_to_console(status)
        elif status[1] == "e":
            add_error_msg_to_console(status)


def load_to_rows_table(filepath: str):
    with open(filepath, 'r') as json_file:
        entries = json.load(json_file)
        print(f'entries {entries}')
        for entry in entries:
            name = entry["name"]
            enabled = entry["enabled"]

            input_fp = entry["input"]["filepath"]
            input_st = entry["input"]["sheetname"]
            selection_mode = selection_to_string(entry["input"]["selection_mode"])

            output_fp = entry["output"]["filepath"]
            output_st = entry["output"]["sheetname"]
            include_name = entry["output"]["include_name"]

            with dpg.table_row(parent="row table"):
                dpg.add_selectable(label=name, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=enabled, span_columns=True, callback=table_row_callback, user_data=entry)

                dpg.add_selectable(label=input_fp, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=input_st, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=selection_mode, span_columns=True, callback=table_row_callback, user_data=entry)

                dpg.add_selectable(label=output_fp, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=output_st, span_columns=True, callback=table_row_callback, user_data=entry)
                dpg.add_selectable(label=include_name, span_columns=True, callback=table_row_callback, user_data=entry)

    json_file.close()


def rows_gui():
    with dpg.child_window(autosize_x=True, height=75):
        dpg.add_input_text(label="Name", tag="Rows Name Input")
        dpg.add_checkbox(label="Enabled", tag="Rows Enabled Input")
    
    with dpg.child_window(autosize_x=True, height=175):
        create_file_input(entry_label="Input Filepath", input_tag="Rows Input FP Input")
        dpg.add_input_text(label="Input Sheetname", tag="Rows Input ST Input")

        dpg.add_radio_button(items=["Row", "Code"], default_value="Row", callback=radio_callback, horizontal=True, tag="Rows Selection Mode Input")

        dpg.add_input_int(label="Row", min_value=1, default_value=1, tag="Rows row option")
        dpg.add_input_text(label="Code", tag="Rows code option")
        dpg.add_input_text(label="Column", tag="Rows column option")

        dpg.hide_item("Rows code option")
        dpg.hide_item("Rows column option")

    with dpg.child_window(autosize_x=True, height=150):
        create_file_input(entry_label="Output Filepath", input_tag="Rows Output FP Input")
        dpg.add_input_text(label="Output Sheetname", tag="Rows Output ST Input")
        dpg.add_checkbox(label="Include Name", tag="Rows Include Name Input")

    with dpg.group(horizontal=True):
        dpg.add_button(label="Add to table", callback=save_to_row_table)
        dpg.add_button(label="Update entry", callback=update_row_entry)
        dpg.add_button(label="Delete entry", callback=delete_row_entry)
        dpg.add_button(label="Save to file", callback=save_to_row_file)
        dpg.add_button(label="Pull data", callback=pull_rows)

    with dpg.table(tag="row table", header_row=True, row_background=True, policy=dpg.mvTable_SizingStretchSame,
                borders_innerH=True, borders_outerH=True, borders_innerV=True,
                borders_outerV=True):
        dpg.add_table_column(label="Name")
        dpg.add_table_column(label="Enabled")
        dpg.add_table_column(label="Input FP")
        dpg.add_table_column(label="Input ST")
        dpg.add_table_column(label="Selection Mode")
        dpg.add_table_column(label="Output FP")
        dpg.add_table_column(label="Output ST")
        dpg.add_table_column(label="Include Name")

    load_to_rows_table("row_project_info.json")